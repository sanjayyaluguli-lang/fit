from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

W,H=13.333,7.5
fx=W/1657; fy=H/951
def X(v): return Inches(v*fx)
def Y(v): return Inches(v*fy)
rgb=lambda h: RGBColor.from_string(h)
NAVY='1D4B6E'; LIME='7FD11B'; DARK='13222D'; GREY='F0F0F0'; GREEN='92D050'; ORANGE='FFC000'; RED='E04646'
prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
s=prs.slides.add_slide(prs.slide_layouts[6])

def box(x,y,w,h,fill=None,line=None,shape=MSO_SHAPE.RECTANGLE,lw=1.5,adj=None):
    sh=s.shapes.add_shape(shape,X(x),Y(y),X(w),Y(h))
    if adj is not None: sh.adjustments[0]=adj
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    else: sh.fill.background()
    if line: sh.line.color.rgb=rgb(line); sh.line.width=Pt(lw)
    else: sh.line.fill.background()
    sh.shadow.inherit=False
    return sh
def text(sh,t,size=10,bold=False,color=DARK,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.MIDDLE,wrap=True):
    tf=sh.text_frame; tf.word_wrap=wrap; tf.vertical_anchor=anchor
    tf.margin_left=tf.margin_right=Inches(.05); tf.margin_top=tf.margin_bottom=Inches(.02)
    lines=t if isinstance(t,list) else [t]
    for i,l in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align
        r=p.add_run(); r.text=l; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color); r.font.name='Arial'
def tb(x,y,w,h,t,**k):
    sh=s.shapes.add_textbox(X(x),Y(y),X(w),Y(h)); text(sh,t,**k); return sh
def panel(x,y,w,h,title):
    box(x,y,w,h,line=NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,adj=.03)
    tb(x,y+8,w,34,title,size=15,bold=True,align=PP_ALIGN.CENTER)
    ln=s.shapes.add_connector(1,X(x+30),Y(y+52),X(x+w-30),Y(y+52)); ln.line.color.rgb=rgb(NAVY); ln.line.width=Pt(1)
def ring(cx,cy,r,val,label=None):
    o=box(cx-r,cy-r,2*r,2*r,line='5B9BD5',shape=MSO_SHAPE.OVAL,lw=7)
    tb(cx-r,cy-14,2*r,28,val,size=8,bold=True,align=PP_ALIGN.CENTER,wrap=False)

# header
p=box(58,16,285,50,fill=NAVY,shape=MSO_SHAPE.ROUNDED_RECTANGLE,adj=.5); text(p,'Project A',20,True,LIME,PP_ALIGN.CENTER)
b=box(1366,20,260,48,fill=LIME,shape=MSO_SHAPE.ROUNDED_RECTANGLE,adj=.15); text(b,'Updated cw{{CW}}',18,True,'1C3A00',PP_ALIGN.CENTER); b.text_frame.word_wrap=False
tb(58,86,1100,56,'SYS VAL PROJECT DASHBOARD REL R250',size=28,bold=True)

# 1 equipment
panel(50,160,370,336,'TEST EQUIPMENT STATUS')
rows=[('Test Platform','Planned Date','Achieved Date','Status'),
('Veh 1','cw43','cw50','Ready'),('Veh 2','cw46/25','cw33/26','Ready'),('SYS.5 Bench','cw40','cw40','Ready'),
('SYS.4 Bench 1','cw37','cw37','Ready'),('CS Bench 1','cw40','cw40','Ready'),('Cluster HIL','cw32/26','cw06/27','Evaluation'),('RadarView','cw45','cw45','Ready')]
t=s.shapes.add_table(len(rows),4,X(60),Y(228),X(350),Y(256)).table
for i,wd in enumerate([110,80,80,80]): t.columns[i].width=X(wd)
for r,row in enumerate(rows):
    t.rows[r].height=Y(256/len(rows))
    for c,v in enumerate(row):
        cell=t.cell(r,c); cell.fill.solid()
        fill='FFFFFF' if r==0 or c<3 else (GREEN if v=='Ready' else ORANGE)
        cell.fill.fore_color.rgb=rgb(fill); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
        cell.margin_left=cell.margin_right=Inches(.03); cell.margin_top=cell.margin_bottom=Inches(.01)
        tf=cell.text_frame; tf.paragraphs[0].alignment=PP_ALIGN.CENTER
        rn=tf.paragraphs[0].add_run(); rn.text=v; rn.font.size=Pt(8); rn.font.name='Arial'; rn.font.bold=(r==0 or c==0 or c==3); rn.font.color.rgb=rgb(DARK)
tb(60,498,350,14,'Edit rows directly in the deck',size=6,color='888888',align=PP_ALIGN.RIGHT)

# 2 coverage
panel(430,158,378,338,'COVERAGE PROGRESS')
tb(510,236,120,22,'SYS.4',size=11,bold=True,align=PP_ALIGN.CENTER); tb(630,236,120,22,'SYS.5',size=11,bold=True,align=PP_ALIGN.CENTER)
ring(571,298,40,'{{SYS4_COVERAGE}}'); ring(685,298,40,'{{SYS5_COVERAGE}}')
tb(450,352,340,24,'Comments :',size=11,bold=True)
tb(450,376,340,60,['1) {{SYS4_COVERAGE}} SYS.4 coverage achieved','2) {{SYS5_COVERAGE}} SYS.5 coverage achieved'],size=9)
tb(450,440,340,48,'{{COVERAGE_NOTE}}',size=9,color='E02020')

# 3 efficiency
panel(818,158,376,338,'TEST EFFICIENCY')
tb(830,236,170,22,'SYS.5 Automation',size=10,bold=True,align=PP_ALIGN.CENTER); tb(1010,236,170,22,'SYS.4 Automation',size=10,bold=True,align=PP_ALIGN.CENTER)
ring(893,290,38,'{{SYS5_AUTO_RATE}}'); ring(1090,290,38,'{{SYS4_AUTO_RATE}}')
tb(830,332,340,20,'Comments :',size=10,bold=True)
tb(830,350,350,56,['New automation target: CW 48 (FFV)','1) SYS.4: {{SYS4_AUTO_COUNT}} auto / {{SYS4_MANUAL_COUNT}} manual','2) SYS.5: {{SYS5_AUTO_COUNT}} auto / {{SYS5_MANUAL_COUNT}} manual'],size=7.5)
tb(830,414,170,22,'Test Duration',size=10,bold=True)
d=box(1010,412,94,24,fill=ORANGE,line='D9A400',shape=MSO_SHAPE.ROUNDED_RECTANGLE,lw=1); text(d,'14 days',10,True,DARK,PP_ALIGN.CENTER)
tb(830,440,350,52,['Comments :','1) Test Duration reduced starting from cw19 (-3d) - done','2) Further improvements to be assessed (cw21) - done'],size=7.5)

# 4 execution
panel(1203,158,424,338,'TEST EXECUTION PROGRESS')
leg=[('TC planned','2E75B6'),('TC Executed','1F9E9E'),('Passed','7FD98A'),('Failed',RED)]
for i,(n,c) in enumerate(leg):
    box(1250+i*88,236,34,6,fill=c); tb(1236+i*88,242,86,16,n,size=7,align=PP_ALIGN.CENTER)
def bar(y,label,keys):
    tb(1208,y,70,40,label,size=12,bold=True,wrap=False)
    for i,((n,c),k) in enumerate(zip(leg,keys)):
        bx=box(1276+i*85,y+2,85,36,fill=c); text(bx,'{{%s}}'%k,6,True,'FFFFFF',PP_ALIGN.CENTER); bx.text_frame.word_wrap=False
bar(262,'SYS.4',['SYS4_TC_PLANNED','SYS4_TC_EXECUTED','SYS4_TC_PASSED','SYS4_TC_FAILED'])
bar(314,'SYS.5',['SYS5_TC_PLANNED','SYS5_TC_EXECUTED','SYS5_TC_PASSED','SYS5_TC_FAILED'])
tb(1220,372,390,20,'Comments :',size=10,bold=True)
tb(1220,392,390,90,['1) Execution started on CW38.1','2) SysBL Audit date 30..09.2026 - YELLOW release'],size=8)

# bug escape
panel(48,518,374,336,'Bug escape rate')
ring(128,656,44,'{{BUG_RATE}}')
tb(66,722,340,22,'Comments :',size=11,bold=True)
tb(66,746,350,104,'{{BUG_COMMENT}}',size=8,anchor=MSO_ANCHOR.TOP)

# support required
sr=box(470,520,264,40,fill=NAVY); text(sr,'SUPPORT REQUIRED',13,True,'FFFFFF',PP_ALIGN.CENTER); sr.text_frame.word_wrap=False
hdr=['N°','ACTION ITEM','RESPONSIBLE','DUE DATE','STATUS']; cw=[68,460,172,220,238]
t=s.shapes.add_table(4,5,X(464),Y(578),X(sum(cw)),Y(306)).table
for i,w in enumerate(cw): t.columns[i].width=X(w)
t.rows[0].height=Y(44)
for r in range(1,4): t.rows[r].height=Y(87)
for r in range(4):
    for c in range(5):
        cell=t.cell(r,c); cell.fill.solid(); cell.vertical_anchor=MSO_ANCHOR.MIDDLE
        if r==0: v=hdr[c]; fill=NAVY; col='FFFFFF'; sz=12
        else:
            v=[str(r),'{{ACTION_%d}}'%r,'{{RESP_%d}}'%r,'{{DUE_%d}}'%r,'{{STATUS_%d}}'%r][c]
            fill=NAVY if c==0 else GREY; col='FFFFFF' if c==0 else DARK; sz=11
        cell.fill.fore_color.rgb=rgb(fill)
        pp=cell.text_frame.paragraphs[0]; pp.alignment=PP_ALIGN.CENTER if (r==0 or c==0) else PP_ALIGN.LEFT
        rn=pp.add_run(); rn.text=v; rn.font.size=Pt(sz); rn.font.name='Arial'; rn.font.color.rgb=rgb(col)

# footer
tb(52,890,420,22,'Valeo   Driving Change Together',size=9,color='555555')
tb(1200,890,430,22,'Strictly private & confidential   |   2',size=9,color=NAVY,align=PP_ALIGN.RIGHT)
prs.save('SYS_VAL_Dashboard_Template.pptx')
